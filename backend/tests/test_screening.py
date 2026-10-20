import hashlib

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.config import Settings
from app.database import Base
from app.models import (
    ApplicationSource,
    Candidate,
    CandidateApplication,
    CVDocument,
    CVProcessingStatus,
    JobDescription,
    JobDescriptionStatus,
    JobPosting,
    JobPostingStatus,
    RankingEntry,
    RankingSnapshot,
    RequirementMatch,
    RequirementProfile,
    RequirementProfileStatus,
    ScreeningBatch,
    ScreeningBatchStatus,
    ScreeningRun,
    ScreeningRunStatus,
    User,
    UserRole,
)
from app.screening.ai import (
    CVProfile,
    EmploymentItem,
    EvidenceItem,
    EvidenceValidation,
    EvidenceValidationItem,
    RequirementAssessment,
    ScreeningAssessment,
    SummaryPoint,
)
from app.screening.parsing import CVParser, CVReviewRequired
from app.services.screening_service import ScreeningService


class FakeScreeningAI:
    model_name = "fake-screening-model"

    def extract_profile(self, *, blocks):
        assert any(block["id"].startswith("docx-paragraph-") for block in blocks)
        return CVProfile(
            candidate_name="Ada Candidate",
            skills=["Python", "FastAPI"],
            employment=[
                EmploymentItem(
                    employer="Example Labs",
                    title="Backend Engineer",
                    start_date="2022-01",
                    end_date="2024-12",
                    responsibilities=["Built Python APIs with FastAPI."],
                    skills=["Python", "FastAPI"],
                    evidence_block_ids=["docx-paragraph-2"],
                )
            ],
        )

    def assess_requirements(self, *, profile, requirements, blocks):
        evidence_block = next(
            block for block in blocks if "Built Python APIs" in block["text"]
        )
        matches = []
        for requirement in requirements:
            if requirement["name"] == "Python":
                matches.append(
                    RequirementAssessment(
                        requirement_key=requirement["key"],
                        status="met",
                        explanation="Professional Python API work is evidenced.",
                        evidence=[
                            EvidenceItem(
                                block_id=evidence_block["id"],
                                quote="Built Python APIs with FastAPI.",
                            )
                        ],
                    )
                )
            else:
                matches.append(
                    RequirementAssessment(
                        requirement_key=requirement["key"],
                        status="unmet",
                        explanation=(
                            f"{requirement['name']} is not evidenced in this CV."
                        ),
                        evidence_gaps=["No supporting CV evidence"],
                    )
                )
        python_key = next(
            item["key"] for item in requirements if item["name"] == "Python"
        )
        docker_key = next(
            item["key"] for item in requirements if item["name"] == "Docker"
        )
        return ScreeningAssessment(
            matches=matches,
            suitability_summary="Ignored in favor of validated summary points.",
            strengths=[
                SummaryPoint(
                    requirement_key=python_key,
                    text="Professional Python API delivery is evidenced.",
                )
            ],
            gaps=[
                SummaryPoint(
                    requirement_key=docker_key,
                    text="Docker is not evidenced in this CV.",
                )
            ],
        )

    def validate_evidence(self, *, requirements, proposed_assessment, blocks):
        return EvidenceValidation(
            validations=[
                EvidenceValidationItem(
                    requirement_key=requirement["key"],
                    supports_conclusion=True,
                    reason="The cited evidence supports the conclusion.",
                )
                for requirement in requirements
            ]
        )


def _settings(storage_root) -> Settings:
    return Settings(
        database_url="sqlite://",
        jwt_secret_key="screening-test-secret-key-that-is-long-enough",
        interviewer_email="interviewer@example.com",
        interviewer_password="StrongInterviewerPassword123!",
        storage_root=str(storage_root),
        cv_min_extract_characters=20,
    )


def _records(db: Session):
    interviewer = User(
        name="Interviewer",
        email="screening-owner@example.com",
        password_hash="hash",
        role=UserRole.INTERVIEWER,
    )
    other = User(
        name="Other",
        email="screening-other@example.com",
        password_hash="hash",
        role=UserRole.CANDIDATE,
    )
    db.add_all([interviewer, other])
    db.flush()
    profile = RequirementProfile(
        title="Backend Engineer",
        required_skills=[{"name": "Python", "weight": 2}],
        preferred_skills=[{"name": "Docker", "weight": 1}],
        qualifications=[],
        selection_criteria=[],
        status=RequirementProfileStatus.CONFIRMED,
        created_by_user_id=interviewer.id,
    )
    db.add(profile)
    db.flush()
    job = JobDescription(
        requirement_profile_id=profile.id,
        title="Backend Engineer",
        content="Backend job",
        structured_content={},
        status=JobDescriptionStatus.APPROVED,
        created_by_user_id=interviewer.id,
    )
    db.add(job)
    db.flush()
    posting = JobPosting(
        job_description_id=job.id,
        platform="discord",
        status=JobPostingStatus.PUBLISHED,
    )
    db.add(posting)
    db.flush()
    return interviewer, other, posting


def _application_with_docx(db, storage_root, posting, index):
    candidate = Candidate(
        name=f"Candidate {index}",
        email=f"candidate{index}@example.com",
    )
    db.add(candidate)
    db.flush()
    application = CandidateApplication(
        job_posting_id=posting.id,
        candidate_id=candidate.id,
        source=ApplicationSource.GOOGLE_FORM,
        source_reference=f"row-{index}",
    )
    db.add(application)
    db.flush()
    relative = f"cvs/{application.id}/candidate-{index}.docx"
    path = storage_root / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    source = Document()
    source.add_paragraph(f"Candidate {index}")
    source.add_paragraph("Built Python APIs with FastAPI.")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, FastAPI"
    source.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    document = CVDocument(
        application_id=application.id,
        candidate_id=candidate.id,
        original_filename=path.name,
        storage_path=relative,
        mime_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        file_size_bytes=path.stat().st_size,
        sha256=digest,
        source=ApplicationSource.GOOGLE_FORM,
        status=CVProcessingStatus.VALIDATED,
    )
    db.add(document)
    return application, document


def test_batch_parses_existing_docx_ranks_ties_and_isolates_missing_cv(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'screening.db'}")
    Base.metadata.create_all(engine)
    settings = _settings(tmp_path / "storage")
    service = ScreeningService(settings, ai=FakeScreeningAI())

    with Session(engine) as db:
        interviewer, other, posting = _records(db)
        first, first_document = _application_with_docx(
            db, tmp_path / "storage", posting, 1
        )
        second, _ = _application_with_docx(db, tmp_path / "storage", posting, 2)
        missing_candidate = Candidate(name="Missing CV")
        db.add(missing_candidate)
        db.flush()
        missing = CandidateApplication(
            job_posting_id=posting.id,
            candidate_id=missing_candidate.id,
            source=ApplicationSource.GOOGLE_FORM,
            source_reference="row-missing",
        )
        db.add(missing)
        db.commit()

        with pytest.raises(LookupError):
            service.create_batch(
                db,
                interviewer_user_id=other.id,
                job_posting_id=posting.id,
                application_ids=None,
            )

        batch = service.create_batch(
            db,
            interviewer_user_id=interviewer.id,
            job_posting_id=posting.id,
            application_ids=None,
        )
        included_ids = set(batch.application_ids)
        assert included_ids == {first.id, second.id, missing.id}

        service.process_batch(db, batch_id=batch.id)
        db.refresh(batch)
        assert batch.status is ScreeningBatchStatus.COMPLETED_WITH_ERRORS
        assert batch.successful_applications == 2
        assert batch.review_applications == 0
        assert batch.failed_applications == 1

        runs = list(db.scalars(select(ScreeningRun).order_by(ScreeningRun.id)))
        assert [run.status for run in runs].count(ScreeningRunStatus.COMPLETED) == 2
        assert [run.status for run in runs].count(
            ScreeningRunStatus.FAILED
        ) == 1
        completed = next(run for run in runs if run.application_id == first.id)
        assert completed.requirement_coverage == pytest.approx(66.67)
        assert completed.met_count == 1
        assert completed.unmet_count == 1

        db.refresh(first_document)
        assert first_document.status is CVProcessingStatus.PROCESSED
        blocks = first_document.extraction_data["source_blocks"]
        assert any(block["id"] == "docx-paragraph-2" for block in blocks)
        assert any(block["id"] == "docx-table-1-row-1" for block in blocks)

        match = db.scalar(
            select(RequirementMatch).where(
                RequirementMatch.screening_run_id == completed.id,
                RequirementMatch.name == "Python",
            )
        )
        assert match.status.value == "met"
        assert match.evidence[0]["quote"] == "Built Python APIs with FastAPI."

        snapshot = db.scalar(
            select(RankingSnapshot).order_by(RankingSnapshot.id.desc())
        )
        entries = list(
            db.scalars(
                select(RankingEntry)
                .where(RankingEntry.ranking_snapshot_id == snapshot.id)
                .order_by(RankingEntry.application_id)
            )
        )
        ranked = [entry for entry in entries if entry.disposition.value == "ranked"]
        review = [entry for entry in entries if entry.disposition.value == "review"]
        assert [entry.rank for entry in ranked] == [1, 1]
        assert len(review) == 0
        assert all(entry.requirement_coverage > 0 for entry in ranked)

        retry_batch = service.create_batch(
            db,
            interviewer_user_id=interviewer.id,
            job_posting_id=posting.id,
            application_ids=None,
        )
        assert retry_batch.application_ids == [missing.id]
    engine.dispose()


def test_screening_schema_exposes_only_three_outcomes_and_normalizes_legacy_status():
    assert set(RequirementAssessment.model_json_schema()["properties"]["status"]["enum"]) == {"met", "partially_met", "unmet"}
    assert RequirementAssessment(requirement_key="skill:0", status="unable_to_determine", explanation="No sufficient evidence").status == "unmet"


def test_unverifiable_positive_evidence_is_downgraded():
    result = ScreeningService._validate_assessment(
        requirements=[
            {
                "key": "required_skill:0",
                "category": "required_skill",
                "name": "Python",
                "details": None,
                "weight": 1.0,
            }
        ],
        blocks=[
            {
                "id": "pdf-page-1",
                "locator": "PDF page 1",
                "text": "Worked with Java services.",
            }
        ],
        assessment={
            "matches": [
                {
                    "requirement_key": "required_skill:0",
                    "status": "met",
                    "explanation": "Python is demonstrated.",
                    "evidence": [
                        {
                            "block_id": "pdf-page-1",
                            "quote": "Built production Python systems.",
                        }
                    ],
                    "evidence_gaps": [],
                }
            ],
            "suitability_summary": "Unsupported summary",
            "strengths": [],
            "gaps": [],
            "points_to_verify": [],
        },
        evidence_validation={
            "validations": [
                {
                    "requirement_key": "required_skill:0",
                    "supports_conclusion": False,
                    "reason": "The CV block describes Java, not Python.",
                }
            ]
        },
    )

    assert result["matches"][0]["status"] == "unmet"
    assert result["requirement_coverage"] == 0
    assert result["suitability_summary"] == ""
    assert "invalid_or_unverifiable_quote" in result["matches"][0]["review_flags"]


def test_weighted_coverage_normalizes_legacy_matches_to_three_outcomes():
    requirements = [
        {
            "key": "required_skill:0",
            "category": "required_skill",
            "name": "Python",
            "details": None,
            "weight": 2.0,
        },
        {
            "key": "preferred_skill:0",
            "category": "preferred_skill",
            "name": "Cloud delivery",
            "details": None,
            "weight": 2.0,
        },
        {
            "key": "qualification:0",
            "category": "qualification",
            "name": "Computer science degree",
            "details": None,
            "weight": 1.0,
        },
        {
            "key": "experience:0",
            "category": "experience",
            "name": "Three years of relevant experience",
            "details": None,
            "weight": 1.0,
        },
    ]
    blocks = [
        {
            "id": "docx-paragraph-1",
            "locator": "DOCX paragraph 1",
            "text": "Built Python APIs and assisted with AWS deployments.",
        }
    ]
    assessment = {
        "matches": [
            {
                "requirement_key": "required_skill:0",
                "status": "met",
                "explanation": "Python delivery is directly evidenced.",
                "evidence": [
                    {
                        "block_id": "docx-paragraph-1",
                        "quote": "Built Python APIs",
                    }
                ],
                "evidence_gaps": [],
            },
            {
                "requirement_key": "preferred_skill:0",
                "status": "partially_met",
                "explanation": "Some AWS deployment exposure is evidenced.",
                "evidence": [
                    {
                        "block_id": "docx-paragraph-1",
                        "quote": "assisted with AWS deployments",
                    }
                ],
                "evidence_gaps": ["Ownership and depth are not evidenced"],
            },
            {
                "requirement_key": "qualification:0",
                "status": "unmet",
                "explanation": "The degree is not evidenced in this CV.",
                "evidence": [],
                "evidence_gaps": ["No education evidence"],
            },
            {
                "requirement_key": "experience:0",
                "status": "unable_to_determine",
                "explanation": "Employment dates are incomplete.",
                "evidence": [],
                "evidence_gaps": ["Missing supported dates"],
            },
        ],
        "strengths": [
            {
                "requirement_key": "required_skill:0",
                "text": "Python API delivery is evidenced.",
            }
        ],
        "gaps": [
            {
                "requirement_key": "qualification:0",
                "text": "The degree is not evidenced in this CV.",
            }
        ],
        "points_to_verify": [
            {
                "requirement_key": "experience:0",
                "text": "Verify employment dates.",
            }
        ],
    }
    evidence_validation = {
        "validations": [
            {
                "requirement_key": requirement["key"],
                "supports_conclusion": True,
                "reason": "The conclusion is appropriately bounded.",
            }
            for requirement in requirements
        ]
    }

    result = ScreeningService._validate_assessment(
        requirements=requirements,
        blocks=blocks,
        assessment=assessment,
        evidence_validation=evidence_validation,
    )

    assert [match["status"] for match in result["matches"]] == [
        "met",
        "partially_met",
        "unmet",
        "unmet",
    ]
    assert result["counts"] == {
        "met": 1,
        "partially_met": 1,
        "unmet": 2,
    }
    assert result["requirement_coverage"] == 50.0


def test_pdf_parser_preserves_page_reference(tmp_path):
    storage_root = tmp_path / "storage"
    path = storage_root / "cvs/1/sample.pdf"
    path.parent.mkdir(parents=True)
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    page[NameObject("/Resources")] = resources
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 72 720 Td (Built Python APIs with FastAPI.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as target:
        writer.write(target)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    document = CVDocument(
        application_id=1,
        candidate_id=1,
        original_filename="sample.pdf",
        storage_path="cvs/1/sample.pdf",
        mime_type="application/pdf",
        file_size_bytes=path.stat().st_size,
        sha256=digest,
        source=ApplicationSource.GOOGLE_FORM,
        status=CVProcessingStatus.VALIDATED,
    )

    parsed = CVParser(_settings(storage_root)).parse(document)
    assert parsed.blocks[0].id == "pdf-page-1"
    assert parsed.blocks[0].locator == "PDF page 1"
    assert "Built Python APIs" in parsed.blocks[0].text

    path.write_bytes(b"not a pdf")
    with pytest.raises(CVReviewRequired, match="content is not"):
        CVParser(_settings(storage_root)).parse(document)
