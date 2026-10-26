import hashlib
import re
import zipfile
import time
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.config import Settings
from app.models.candidate import CVDocument
from app.screening.ocr import LocalCVOCR, OCRReadError, ocr_page_notice


SUPPORTED_SUFFIXES = {".pdf", ".docx"}


class CVReviewRequired(ValueError):
    """The document cannot receive a reliable automatic assessment."""


@dataclass(frozen=True)
class SourceBlock:
    id: str
    locator: str
    text: str

    def as_dict(self) -> dict[str, str]:
        return {"id": self.id, "locator": self.locator, "text": self.text}


@dataclass(frozen=True)
class ParsedCV:
    text: str
    blocks: list[SourceBlock]
    parser: str
    page_count: int | None
    warnings: list[str]
    ocr: dict = field(default_factory=dict)


class CVParser:
    def __init__(self, settings: Settings, *, ocr=None):
        self.settings = settings
        self.ocr = ocr or LocalCVOCR(settings)

    def parse(self, document: CVDocument) -> ParsedCV:
        path = self._trusted_path(document.storage_path)
        self._validate_file(path, document)
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            parsed = self._parse_pdf(path)
        elif suffix == ".docx":
            parsed = self._parse_docx(path)
        else:
            raise CVReviewRequired("Only PDF and DOCX CV files are supported")
        normalized = parsed.text.strip()
        if len(normalized) < self.settings.cv_min_extract_characters:
            raise CVReviewRequired(
                "The CV contains too little extractable text. It may be scanned, "
                "encrypted, empty, or corrupt; upload a readable text-based replacement."
            )
        if len(normalized) > self.settings.cv_max_characters:
            raise CVReviewRequired(
                "The CV exceeds the configured extraction complexity limit; "
                "upload a shorter CV instead of silently truncating it."
            )
        return parsed

    def _trusted_path(self, relative_path: str) -> Path:
        root = Path(self.settings.storage_root).expanduser().resolve()
        cv_root = (root / "cvs").resolve()
        path = (root / relative_path).resolve()
        if path != cv_root and cv_root not in path.parents:
            raise CVReviewRequired("The saved CV reference is outside private storage")
        if not path.is_file():
            raise CVReviewRequired("The saved CV file is missing from private storage")
        return path

    def _validate_file(self, path: Path, document: CVDocument) -> None:
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise CVReviewRequired("Only PDF and DOCX CV files are supported")
        size = path.stat().st_size
        max_bytes = self.settings.cv_max_file_size_mb * 1024 * 1024
        if size == 0:
            raise CVReviewRequired("The saved CV file is empty")
        if size > max_bytes:
            raise CVReviewRequired(
                f"The CV exceeds the {self.settings.cv_max_file_size_mb} MB limit"
            )
        if suffix == ".pdf":
            with path.open("rb") as source:
                if source.read(5) != b"%PDF-":
                    raise CVReviewRequired("The file extension is PDF but its content is not")
        elif not zipfile.is_zipfile(path):
            raise CVReviewRequired("The file extension is DOCX but its content is not")

        digest = hashlib.sha256()
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
        if document.sha256 and digest.hexdigest() != document.sha256:
            raise CVReviewRequired(
                "The locally stored CV no longer matches its recorded checksum"
            )

    def _parse_pdf(self, path: Path) -> ParsedCV:
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise CVReviewRequired("Encrypted PDFs cannot be processed; upload an unencrypted replacement")
            if len(reader.pages) > self.settings.cv_max_pages:
                raise CVReviewRequired(
                    f"The PDF exceeds the {self.settings.cv_max_pages}-page limit"
                )
            blocks: list[SourceBlock] = []
            warnings: list[str] = []
            ocr_pages = []
            deadline = time.monotonic() + self.settings.cv_ocr_timeout_seconds
            for page_number, page in enumerate(reader.pages, start=1):
                text = self._clean_text(page.extract_text() or "")
                locator = f"PDF page {page_number}"
                # Sparse selectable headers may cover a scanned body. OCR replaces
                # the whole page, never concatenates duplicate native/OCR text.
                needs_ocr = not text or (len(text) < self.settings.cv_min_extract_characters and bool(page.images))
                if not text and not page.images and hasattr(page, "get_contents") and not page.get_contents():
                    warnings.append(f"PDF page {page_number} is blank")
                    continue
                if needs_ocr:
                    if len(ocr_pages) >= self.settings.cv_ocr_max_pages:
                        raise CVReviewRequired("CV OCR exceeds the scanned-page limit. Upload a shorter or text-based PDF/DOCX; no candidate score was assigned.")
                    text, metadata = self.ocr.read_page(path, page_number, deadline=deadline)
                    text = self._clean_text(text)
                    ocr_pages.append(metadata)
                    if text:
                        locator += " (OCR)"
                        warnings.append(ocr_page_notice(page_number))
                if text:
                    blocks.append(
                        SourceBlock(
                            id=f"pdf-page-{page_number}",
                            locator=locator,
                            text=text,
                        )
                    )
                else:
                    warnings.append(f"No text was extracted from PDF page {page_number}")
                if sum(len(block.text) for block in blocks) > self.settings.cv_max_characters:
                    raise CVReviewRequired("The CV exceeds the configured extraction complexity limit; upload a shorter CV.")
        except OCRReadError as exc:
            raise CVReviewRequired(str(exc)) from None
        except CVReviewRequired:
            raise
        except Exception as exc:
            raise CVReviewRequired("The PDF is corrupt or unreadable") from exc
        return ParsedCV(
            text="\n\n".join(block.text for block in blocks),
            blocks=blocks,
            parser="pypdf+tesseract" if ocr_pages else "pypdf",
            page_count=len(reader.pages),
            warnings=warnings,
            ocr={"pages": ocr_pages} if ocr_pages else {},
        )

    def _parse_docx(self, path: Path) -> ParsedCV:
        try:
            source = Document(str(path))
            blocks: list[SourceBlock] = []
            for index, paragraph in enumerate(source.paragraphs, start=1):
                text = self._clean_text(paragraph.text)
                if text:
                    blocks.append(
                        SourceBlock(
                            id=f"docx-paragraph-{index}",
                            locator=f"DOCX paragraph {index}",
                            text=text,
                        )
                    )
            for table_index, table in enumerate(source.tables, start=1):
                for row_index, row in enumerate(table.rows, start=1):
                    values = [self._clean_text(cell.text) for cell in row.cells]
                    text = " | ".join(value for value in values if value)
                    if text:
                        blocks.append(
                            SourceBlock(
                                id=f"docx-table-{table_index}-row-{row_index}",
                                locator=(
                                    f"DOCX table {table_index}, row {row_index}"
                                ),
                                text=text,
                            )
                        )
        except Exception as exc:
            raise CVReviewRequired("The DOCX file is corrupt or unreadable") from exc
        return ParsedCV(
            text="\n\n".join(block.text for block in blocks),
            blocks=blocks,
            parser="python-docx",
            page_count=None,
            warnings=[],
        )

    @staticmethod
    def _clean_text(value: str) -> str:
        value = value.replace("\x00", " ")
        value = re.sub(r"[ \t]+", " ", value)
        value = re.sub(r"\n{3,}", "\n\n", value)
        return value.strip()
